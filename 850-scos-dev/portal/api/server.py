"""Portal HTTP Server — serves PWA + REST API. Reads from cache, zero business logic."""
import os, sys, io, json, hashlib, secrets, sqlite3
from datetime import datetime, timedelta
from http.server import HTTPServer, SimpleHTTPRequestHandler
from socketserver import ThreadingMixIn
import urllib.parse, re

PORT = int(os.environ.get('SCOS_PORTAL_PORT', '5050'))  # dev: 5051
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DASHBOARD = os.path.join(ROOT, 'dashboard')
RECEIVER = os.path.join(ROOT, 'receiver')
# DB paths — all inside 850-scos/data/, configurable via env vars
USER_DB = os.environ.get('SCOS_USER_DB') or os.path.join(ROOT, '..', 'data', 'users.db')
PORTAL_DB = os.environ.get('SCOS_PORTAL_DB') or os.path.join(ROOT, '..', 'data', 'portal.db')
SESSION_DAYS = 30
_sessions = {}

if sys.platform == 'win32': sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
os.makedirs(DASHBOARD, exist_ok=True)
os.makedirs(os.path.dirname(USER_DB), exist_ok=True)
# Init user table if not exists
try:
    db = sqlite3.connect(USER_DB)
    db.execute("CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, status TEXT DEFAULT 'pending', created_at TEXT DEFAULT (datetime('now','localtime')))")
    db.commit(); db.close()
except: pass


RECEIVER_SCHEMA = '''
CREATE TABLE IF NOT EXISTS orders (
    po TEXT NOT NULL, po_line TEXT NOT NULL DEFAULT '1',
    region TEXT, sub_type TEXT, priority TEXT, cto_p1 TEXT,
    mcid TEXT, ship_mode TEXT, scac TEXT, master_type TEXT, cust TEXT, ship_status TEXT,
    project TEXT DEFAULT 'DAISY',
    po_qty INTEGER DEFAULT 0, remain_qty INTEGER DEFAULT 0, ship_qty INTEGER DEFAULT 0,
    msbd TEXT, psd TEXT, final_msbd TEXT, po_received TEXT,
    status TEXT, ack_status TEXT, is_hold TEXT, hold_code TEXT, status_label TEXT,
    asn TEXT, hawb TEXT, dpn TEXT, ipn TEXT, dell_so TEXT,
    ship_to_country TEXT, description TEXT,
    input_cdt TEXT, stockin_cdt TEXT, sn_cdt TEXT, createasn_cdt TEXT,
    stbl INTEGER DEFAULT 0, atb INTEGER DEFAULT 0, wip INTEGER DEFAULT 0,
    fg INTEGER DEFAULT 0, sn INTEGER DEFAULT 0,
    actual_shipped INTEGER DEFAULT 0, shipped_qty INTEGER DEFAULT 0, asn_pending INTEGER DEFAULT 0,
    updated_at TEXT, hash TEXT,
    PRIMARY KEY (po, po_line)
);
CREATE TABLE IF NOT EXISTS cache (key TEXT PRIMARY KEY, data TEXT NOT NULL, computed_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS sync_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT, sync_time TEXT,
    added INTEGER DEFAULT 0, updated INTEGER DEFAULT 0, total INTEGER DEFAULT 0
);
'''

def _db():
    conn = sqlite3.connect(PORTAL_DB)
    conn.row_factory = sqlite3.Row
    conn.executescript(RECEIVER_SCHEMA)
    try: conn.execute('ALTER TABLE orders ADD COLUMN shipped_qty INTEGER DEFAULT 0')
    except sqlite3.OperationalError: pass
    try: conn.execute('ALTER TABLE orders ADD COLUMN asn_pending INTEGER DEFAULT 0')
    except sqlite3.OperationalError: pass
    try: conn.execute('ALTER TABLE orders ADD COLUMN cust TEXT')
    except sqlite3.OperationalError: pass
    try: conn.execute('ALTER TABLE orders ADD COLUMN ship_status TEXT')
    except sqlite3.OperationalError: pass
    try: conn.execute('ALTER TABLE orders ADD COLUMN project TEXT DEFAULT \'DAISY\'')
    except sqlite3.OperationalError: pass
    conn.commit()
    return conn

def _hash_pw(pw): return hashlib.sha256(pw.encode()).hexdigest()

def _auth(handler):
    c = handler.headers.get('Cookie', '')
    for part in c.split(';'):
        if part.strip().startswith('ym_token='):
            token = part.strip()[9:]
            if token in _sessions and datetime.now() < _sessions[token]['expires']:
                db = sqlite3.connect(USER_DB)
                row = db.execute('SELECT status FROM users WHERE username=?', (_sessions[token]['user'],)).fetchone()
                db.close()
                if row and row[0] == 'active': return _sessions[token]['user']
            del _sessions[token]
    return None

def _require(handler):
    u = _auth(handler)
    if not u:
        handler.send_response(302); handler.send_header('Location', '/login.html'); handler.end_headers()
    return u

class ThreadingHTTPServer(ThreadingMixIn, HTTPServer): daemon_threads = True


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=DASHBOARD, **kwargs)

    def do_GET(self):
        p = urllib.parse.urlparse(self.path).path
        if p in ('/', '/index.html', '/k1.html'):
            if not _require(self): return
        try:
            if p == '/api/logout':
                self._logout()
            elif p == '/api/health':
                self._health()
            # Phase 2 — standardized Service Layer endpoints（Portal/AI 唯一业务入口）
            elif p == '/api/k1':
                self._serve_cache('k1_summary')
            elif p == '/api/daily':
                self._serve_cache('daily_summary')
            elif p == '/api/kpi':
                self._serve_cache('kpi')
            elif p == '/api/e2e':
                self._serve_cache('e2e_kpi')
            elif p == '/api/risk':
                self._serve_cache('risks')
            elif p == '/api/risk-summary':
                self._serve_cache('risk_summary')
            elif p == '/api/nack':
                self._query_nack()
            elif p == '/api/asn':
                self._query_asn()
            elif p == '/api/st':
                self._query_st()
            elif p == '/api/cto-asn-missing':
                self._serve_cache('cto_asn_missing')
            # compat: old dashboard endpoints
            elif p == '/api/k1-summary':
                self._serve_cache('k1_summary')
            elif p == '/api/daily-summary':
                self._serve_cache('daily_summary')
            elif p == '/api/sort-data':      # compat: old dashboard drill-down
                self._query_orders()
            elif p == '/api/shipped-history':  # compat
                self._query_orders()
            elif p.startswith('/api/cache/'):
                key = p.split('/')[-1]
                self._serve_cache(key)
            elif p == '/api/admin-users':
                self._admin_users()
            elif p == '/api/orders':
                self._query_orders()
            elif p == '/api/export/orders':
                self._export_orders()
            elif p in ('/sw.js', '/manifest.json'):
                self._serve_static(p)
            else:
                if p in ('/', ''):
                    self.path = '/index.html'
                    p = self.path
                if p.endswith('.html'):
                    self._serve_html_injected(p)
                else:
                    super().do_GET()
        except Exception as e:
            self._json({'error': str(e)}, 500)

    def do_POST(self):
        p = urllib.parse.urlparse(self.path).path
        try:
            if p == '/sync':
                self._sync()
            elif p == '/api/login': self._login()
            elif p == '/api/register': self._register()
            elif p == '/api/admin-login': self._admin_login()
            elif p == '/api/admin-action': self._admin_action()
            elif p == '/api/admin-users': self._admin_users()
            else: self._json({'error': 'Not found'}, 404)
        except Exception as e:
            self._json({'error': str(e)}, 500)

    def _health(self):
        conn = _db()
        total = conn.execute('SELECT COUNT(*) as c FROM orders').fetchone()['c']
        last = conn.execute('SELECT MAX(sync_time) as t FROM sync_log').fetchone()
        has_k1 = conn.execute("SELECT 1 FROM cache WHERE key='k1_summary'").fetchone()
        has_risks = conn.execute("SELECT 1 FROM cache WHERE key='risks'").fetchone()
        has_kpi = conn.execute("SELECT 1 FROM cache WHERE key='kpi'").fetchone()
        conn.close()
        self._json({'status': 'ok', 'server_time': datetime.now().isoformat(),
            'db_records': total, 'last_sync_time': last['t'] if last else None,
            'k1_cached': bool(has_k1), 'risks_cached': bool(has_risks),
            'kpi_cached': bool(has_kpi), 'version': 'scos-1.7-dev'})

    AI_WIDGET = '''
<style>
#proj-bar{position:fixed;top:12px;right:16px;z-index:9997;display:flex;gap:6px;font-family:system-ui}
#proj-bar button{padding:6px 14px;border-radius:20px;border:1px solid #334155;background:transparent;color:#94a3b8;font-size:11px;font-weight:700;cursor:pointer}
#proj-bar button.on{background:#3b82f6;border-color:#3b82f6;color:#fff}
</style>
<div id="proj-bar">
  <button data-p="DAISY" onclick="switchProj(this)">DAISY</button>
  <button data-p="SWAN" onclick="switchProj(this)">SWAN</button>
</div>
<script>
function switchProj(btn){var p=btn.getAttribute('data-p');document.cookie='scos_project='+p+';path=/;max-age=31536000';location.reload();}
(function(){var m=document.cookie.match(/scos_project=([A-Za-z0-9_]+)/);var cur=m?m[1]:'DAISY';
var bs=document.querySelectorAll('#proj-bar button');for(var i=0;i<bs.length;i++){if(bs[i].getAttribute('data-p')===cur)bs[i].className='on';}})();
</script>
<style>
#ai-fab{position:fixed;right:18px;bottom:18px;z-index:9999;width:54px;height:54px;border-radius:50%;background:linear-gradient(135deg,#7aa2f7,#bb9af7);border:none;cursor:pointer;font-size:24px;box-shadow:0 6px 20px rgba(0,0,0,.5);display:flex;align-items:center;justify-content:center;transition:transform .15s}
#ai-fab:hover{transform:scale(1.08)}
#ai-drawer{position:fixed;right:14px;bottom:84px;z-index:9998;width:420px;max-width:94vw;height:560px;max-height:74vh;border-radius:14px;border:1px solid #334155;box-shadow:0 12px 40px rgba(0,0,0,.6);display:none;overflow:hidden;background:#0f172a}
#ai-drawer iframe{width:100%;height:100%;border:none;display:none}
#ai-hint{display:none;padding:40px 20px;text-align:center;color:#94a3b8;font-size:13px;line-height:2;font-family:system-ui}
</style>
<button id="ai-fab" onclick="aiToggle()" title="AI Assistant">🤖</button>
<div id="ai-drawer">
  <iframe id="ai-frame" src="about:blank"></iframe>
  <div id="ai-hint">AI 服务未启动<br>双击 <b>start-ai.bat</b> 后刷新页面</div>
</div>
<script>
function aiToggle(){
  var d=document.getElementById('ai-drawer'),f=document.getElementById('ai-frame'),h=document.getElementById('ai-hint');
  var open=d.style.display!=='block';
  d.style.display=open?'block':'none';
  if(open&&(f.src==='about:blank'||!f.src)){
    fetch('http://localhost:5099/',{method:'GET'}).then(function(){
      h.style.display='none';f.style.display='block';f.src='http://localhost:5099/';
    }).catch(function(){h.style.display='block';f.style.display='none';});
  }
}
</script>
'''

    def _serve_html_injected(self, path):
        """静态 HTML + AI 悬浮球注入（所有 dashboard 页面统一获得悬浮对话入口）。"""
        fp = os.path.join(DASHBOARD, path.lstrip('/'))
        if not os.path.isfile(fp):
            self.send_error(404)
            return
        with open(fp, encoding='utf-8', errors='replace') as f:
            content = f.read()
        content = content.replace('</body>', self.AI_WIDGET + '</body>')
        data = content.encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _proj(self):
        """当前项目（scos_project cookie），默认 DAISY（向后兼容）。"""
        c = self.headers.get('Cookie', '')
        m = re.search(r'scos_project=([A-Za-z0-9_]+)', c)
        return m.group(1) if m else 'DAISY'

    def _pkey(self, key):
        """缓存 key 命名空间：DAISY 用原名（旧数据零迁移），其他项目加前缀。"""
        p = self._proj()
        return f'{p}:{key}' if p != 'DAISY' else key

    def _serve_cache(self, key):
        pkey = self._pkey(key)
        conn = _db()
        row = conn.execute('SELECT data, computed_at FROM cache WHERE key=?', (pkey,)).fetchone()
        conn.close()
        if not row:
            self._json({'error': f'No cache for {key}'}, 404)
            return
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        if qs.get('meta', [''])[0] == '1':
            # Phase 2: data provenance for AI Engine — 数据来源 + 更新时间
            self._json({'key': pkey, 'computed_at': row['computed_at'], 'data': json.loads(row['data'])})
        else:
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(row['data'].encode('utf-8'))

    def _build_orders_sql(self, qs):
        """过滤条件构建 — /api/orders 与 /api/export/orders 共用同一套逻辑。项目过滤随 cookie。"""
        sql, params = 'SELECT * FROM orders WHERE project=? AND 1=1', [self._proj()]
        for f in ('region','sub_type','ack_status','ship_mode','scac','mcid','cto_p1','priority','is_hold','cust'):
            v = (qs.get(f, [''])[0]).strip()
            if v: sql += f' AND {f}=?'; params.append(v)
        shipped = qs.get('shipped', [''])[0]
        if shipped == '1': sql += ' AND actual_shipped=1'
        elif shipped == '0': sql += ' AND actual_shipped=0'
        # Date filters (for drill-down from K1 dashboard)
        for df, col in [('msbd','msbd'),('sn','sn_cdt'),('po_received','po_received')]:
            v = (qs.get(df, [''])[0]).strip()
            if v: sql += f' AND {col} LIKE ?'; params.append(v[:10]+'%')
        for df, col in [('po_from','po_received'),('po_to','po_received')]:
            v = (qs.get(df, [''])[0]).strip()
            if v and df == 'po_from': sql += f' AND {col} >= ?'; params.append(v)
            elif v and df == 'po_to': sql += f' AND {col} <= ?'; params.append(v)
        # MSBD range filter (Excel Generator)
        for df, col in [('msbd_from','msbd'),('msbd_to','msbd')]:
            v = (qs.get(df, [''])[0]).strip()
            if v and df == 'msbd_from': sql += f' AND {col} >= ?'; params.append(v)
            elif v and df == 'msbd_to': sql += f' AND {col} <= ?'; params.append(v)
        # GPP status filter
        gpp = (qs.get('gpp', [''])[0]).strip().upper()
        if gpp in ('STBL','ATB','WIP','FG','SN'):
            sql += f' AND {gpp.lower()}>0'
        # CTO 28H exact deadline filter: PO_RECEIVE + 28H falls on this date
        cto28 = (qs.get('cto_28h', [''])[0]).strip()
        if cto28:
            from datetime import timedelta
            try:
                dl = datetime.strptime(cto28[:10], '%Y-%m-%d')
                lo = (dl - timedelta(hours=28)).strftime('%Y-%m-%dT%H:%M:%S')
                hi = dl.strftime('%Y-%m-%dT%H:%M:%S')
                sql += ' AND po_received >= ? AND po_received < ?'
                params.append(lo); params.append(hi)
            except: pass
        limit = qs.get('limit', [''])[0]
        if limit: sql += ' LIMIT ?'; params.append(min(int(limit), 50000))
        return sql, params

    COL_MAP = {'po':'PO','po_line':'PO_LINE','region':'REGION','sub_type':'SUB_TYPE',
        'priority':'PRIORITY','cto_p1':'CTO_P1','mcid':'MCID','ship_mode':'SHIP_MODE',
        'scac':'SCAC','master_type':'MASTER_TYPE','cust':'CUST','ship_status':'SHIP_STATUS','po_qty':'PO_QTY','remain_qty':'REMAIN_QTY',
        'ship_qty':'SHIP_QTY','msbd':'MSBD','psd':'PSD','final_msbd':'FINAL_MSBD',
        'po_received':'PO_RECEIVE_DATE','status':'STATUS','ack_status':'ACK_STATUS',
        'is_hold':'IS_HOLD','hold_code':'HOLD_CODE','status_label':'STATUS_LABEL',
        'asn':'ASN','hawb':'HAWB','dpn':'DPN','ipn':'IPN','dell_so':'DELL_SO',
        'ship_to_country':'SHIP_TO_COUNTRY','description':'DESCRIPTION',
        'input_cdt':'INPUT_CDT','stockin_cdt':'STOCKIN_CDT','sn_cdt':'SN_CDT',
        'createasn_cdt':'CREATEASN_CDT','stbl':'STBL','atb':'ATB','wip':'WIP','fg':'FG','sn':'SN',
        'actual_shipped':'ACTUAL_SHIPPED','shipped_qty':'SHIPPED_QTY'}

    def _map_order_row(self, r):
        return {new_k: r[old_k] for old_k, new_k in self.COL_MAP.items() if old_k in r}

    def _query_orders(self):
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        conn = _db()
        sql, params = self._build_orders_sql(qs)
        rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
        conn.close()
        mapped = [self._map_order_row(r) for r in rows]
        self._json({'records': mapped, 'total': len(mapped)})

    def _export_orders(self):
        """Excel Generator — 过滤 → openpyxl → XLSX 下载。零业务逻辑（纯数据导出）。"""
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        conn = _db()
        sql, params = self._build_orders_sql(qs)
        rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
        conn.close()
        mapped = [self._map_order_row(r) for r in rows]

        import io as _io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = 'orders'
        if mapped:
            headers = list(mapped[0].keys())
            ws.append(headers)
            hfont = Font(bold=True, color='FFFFFF')
            hfill = PatternFill('solid', fgColor='3B82F6')
            for c in ws[1]:
                c.font = hfont
                c.fill = hfill
            for m in mapped:
                ws.append([m.get(h, '') for h in headers])
            ws.freeze_panes = 'A2'
            for i, h in enumerate(headers, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(10, min(32, len(str(h)) + 2))
        else:
            ws.append(['PO'])

        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f'850_orders_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        self.send_response(200)
        self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.send_header('Content-Disposition', f'attachment; filename="{fname}"')
        self.send_header('Content-Length', str(len(buf.getvalue())))
        self.end_headers()
        self.wfile.write(buf.getvalue())

    def _query_st(self):
        """ST（SHIP_STATUS）校验查询 — 读 st_check 缓存（Business Engine 计算结果）。零业务逻辑。"""
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        po = (qs.get('po', [''])[0]).strip()
        asn = (qs.get('asn', [''])[0]).strip()
        if not po and not asn:
            self._json({'error': 'Missing po or asn param'}, 400)
            return
        conn = _db()
        row = conn.execute('SELECT data, computed_at FROM cache WHERE key=?', (self._pkey('st_check'),)).fetchone()
        conn.close()
        if not row:
            self._json({'error': 'st_check cache not available', 'computed_at': None}, 404)
            return
        data = json.loads(row['data'])
        entry = None
        if po:
            entry = data.get('by_po', {}).get(po)
        if entry is None and asn:
            entry = data.get('by_asn', {}).get(asn)
        if entry is None:
            self._json({'found': False, 'query': po or asn, 'verdict': 'NOT_FOUND',
                        'computed_at': row['computed_at'], 'meta': data.get('__meta', {})}, 404)
            return
        self._json({'found': True, 'computed_at': row['computed_at'],
                    'meta': data.get('__meta', {}), 'check': entry})

    def _query_nack(self):
        """NACK 订单查询 — 复用 _query_orders 的过滤/映射，固定 ack_status=REJECT。零业务逻辑。"""
        parsed = urllib.parse.urlparse(self.path)
        extra = parsed.query
        self.path = parsed.path + '?ack_status=REJECT' + (('&' + extra) if extra else '')
        self._query_orders()

    def _query_asn(self):
        """ASN 校验查询 — 读 asn_check 缓存（Business Engine 计算结果）。零业务逻辑。"""
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        asn = (qs.get('asn', [''])[0]).strip()
        if not asn:
            self._json({'error': 'Missing asn param'}, 400)
            return
        conn = _db()
        row = conn.execute('SELECT data, computed_at FROM cache WHERE key=?', (self._pkey('asn_check'),)).fetchone()
        conn.close()
        if not row:
            self._json({'error': 'asn_check cache not available', 'computed_at': None}, 404)
            return
        data = json.loads(row['data'])
        entry = data.get(asn)
        if entry is None:
            self._json({'found': False, 'asn': asn, 'result': 'NOT_FOUND',
                        'computed_at': row['computed_at'], 'meta': data.get('__meta', {})}, 404)
            return
        self._json({'found': True, 'computed_at': row['computed_at'],
                    'meta': data.get('__meta', {}), 'check': entry})

    def _sync(self):
        sys.path.insert(0, RECEIVER)
        from sync import handle_sync
        length = int(self.headers.get('Content-Length', 0))
        raw = self.rfile.read(length) if length > 0 else b''
        if self.headers.get('Content-Encoding', '').lower() == 'gzip':
            import gzip
            raw = gzip.decompress(raw)
        body = json.loads(raw) if raw else {}
        self._json(handle_sync(body))

    def _login(self):
        length = int(self.headers.get('Content-Length', 0))
        body = json.loads(self.rfile.read(length)) if length > 0 else {}
        u, p = body.get('username','').strip(), body.get('password','').strip()
        db = sqlite3.connect(USER_DB)
        row = db.execute('SELECT password_hash,status FROM users WHERE username=?', (u,)).fetchone()
        db.close()
        if not row: self._json({'error': 'Account not found'}, 401)
        elif row[0] != _hash_pw(p): self._json({'error': 'Wrong password'}, 401)
        elif row[1] == 'pending': self._json({'error': 'Account pending approval'}, 403)
        elif row[1] == 'disabled': self._json({'error': 'Account disabled'}, 403)
        else:
            token = secrets.token_hex(32)
            _sessions[token] = {'user': u, 'expires': datetime.now() + timedelta(days=SESSION_DAYS)}
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Set-Cookie', f'ym_token={token}; Path=/; Max-Age={SESSION_DAYS*86400}; HttpOnly; SameSite=Lax')
            self.end_headers()
            self.wfile.write(json.dumps({'ok': True, 'user': u}).encode('utf-8'))

    def _register(self):
        length = int(self.headers.get('Content-Length', 0))
        body = json.loads(self.rfile.read(length)) if length > 0 else {}
        u, p = body.get('username','').strip(), body.get('password','').strip()
        if len(u) < 2 or len(p) < 4: self._json({'error': 'Username >=2 chars, password >=4 chars'}, 400); return
        try:
            db = sqlite3.connect(USER_DB)
            db.execute('INSERT INTO users (username,password_hash) VALUES (?,?)', (u, _hash_pw(p)))
            db.commit(); db.close()
            self._json({'ok': True, 'msg': 'Submitted, awaiting approval'})
        except sqlite3.IntegrityError: self._json({'error': 'Account already exists'}, 400)

    ADMIN_PW = os.environ.get('SCOS_ADMIN_PW', 'admin850')
    _admin_tokens = set()

    def _admin_login(self):
        length = int(self.headers.get('Content-Length', 0))
        body = json.loads(self.rfile.read(length)) if length > 0 else {}
        if body.get('password', '') == self.ADMIN_PW:
            t = secrets.token_hex(16)
            Handler._admin_tokens.add(t)
            self._json({'ok': True, 'token': t})
        else:
            self._json({'error': 'Wrong admin password'}, 403)

    def _check_admin(self):
        qs = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        token = qs.get('token', [''])[0]
        if not token:
            try:
                length = int(self.headers.get('Content-Length', 0))
                if length > 0:
                    body = json.loads(self.rfile.read(length))
                    token = body.get('token', '')
            except: pass
        return token in Handler._admin_tokens if token else False

    def _admin_users(self):
        if not self._check_admin(): self._json({'error': 'Unauthorized'}, 403); return
        db = sqlite3.connect(USER_DB)
        rows = db.execute('SELECT username,status,created_at FROM users ORDER BY created_at DESC').fetchall()
        db.close()
        self._json({'users': [{'username': r[0], 'status': r[1], 'created': r[2]} for r in rows]})

    def _admin_action(self):
        length = int(self.headers.get('Content-Length', 0))
        body = json.loads(self.rfile.read(length)) if length > 0 else {}
        token = body.get('token', '')
        if token not in Handler._admin_tokens: self._json({'error': 'Unauthorized'}, 403); return
        act = body.get('action', '')
        username = body.get('username', '')
        db = sqlite3.connect(USER_DB)
        if act == 'approve':
            db.execute('UPDATE users SET status=? WHERE username=?', ('active', username))
        elif act == 'disable':
            db.execute('UPDATE users SET status=? WHERE username=?', ('disabled', username))
        elif act == 'enable':
            db.execute('UPDATE users SET status=? WHERE username=?', ('active', username))
        elif act == 'approve-all':
            db.execute("UPDATE users SET status='active' WHERE status='pending'")
        db.commit(); db.close()
        self._json({'ok': True})

    def _logout(self):
        c = self.headers.get('Cookie', '')
        for part in c.split(';'):
            if part.strip().startswith('ym_token='):
                _sessions.pop(part.strip()[9:], None)
        self.send_response(302); self.send_header('Location', '/login.html')
        self.send_header('Set-Cookie', 'ym_token=; Path=/; Max-Age=0'); self.end_headers()

    def _serve_static(self, path):
        fp = os.path.join(DASHBOARD, path.lstrip('/'))
        if os.path.exists(fp):
            ct = 'application/javascript' if path.endswith('.js') else 'application/json'
            cc = 'no-cache' if path.endswith('.js') else 'public, max-age=86400'
            self.send_response(200)
            self.send_header('Content-Type', f'{ct}; charset=utf-8')
            self.send_header('Cache-Control', cc)
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(open(fp, 'rb').read())

    def _json(self, data, code=200):
        self.send_response(code)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False, default=str).encode('utf-8'))

    def log_message(self, *args): pass


def main():
    server = ThreadingHTTPServer(('0.0.0.0', PORT), Handler)
    print(f'\n  850 SCOS Portal\n  http://localhost:{PORT}\n  DB: {PORTAL_DB}\n  Ready.\n')
    server.serve_forever()

if __name__ == '__main__':
    main()
